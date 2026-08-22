
"""任务产物在线预览端点测试。

覆盖：
- 未登录预览返回 401
- 预览他人任务返回 404
- PDF / 图片返回 inline 文件流（Content-Disposition: inline）
- XLSX / DOCX 渲染为 HTML（含表格 / 段落内容）
- ZIP 返回 400（不支持预览）
"""

from __future__ import annotations

import io
import sys
import uuid
from pathlib import Path

_API_DIR = Path(__file__).resolve().parent.parent
if str(_API_DIR) not in sys.path:
    sys.path.insert(0, str(_API_DIR))

import pytest
from fastapi.testclient import TestClient

from app.config import get_settings
from app.main import app
from app.models.enums import ConversionType, TaskStatus
from app.state import get_task_manager
from app.utils.file_utils import generate_task_id

_TMP_ROOT = Path(__file__).resolve().parent.parent.parent / "logs" / "test-tmp"


@pytest.fixture()
def client(monkeypatch):
    """每个测试使用独立临时数据库 / 上传 / 输出目录。"""
    _TMP_ROOT.mkdir(parents=True, exist_ok=True)
    base = Path(_TMP_ROOT) / f"prev_{uuid.uuid4().hex[:10]}"
    base.mkdir(parents=True, exist_ok=True)
    settings = get_settings()
    monkeypatch.setattr(settings, "db_path", base / "test.db")
    monkeypatch.setattr(settings, "upload_dir", base / "uploads")
    monkeypatch.setattr(settings, "output_dir", base / "outputs")
    with TestClient(app) as c:
        yield c


def _register_and_login(client, username: str, password: str) -> dict:
    resp = client.post(
        "/api/v1/auth/register",
        json={"username": username, "password": password},
    )
    assert resp.status_code == 200, resp.text
    resp = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": password},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    return {
        "Authorization": "Bearer " + data["access_token"],
        "user_id": int(data["user"]["id"]),
    }


def _make_task_with_files(client, headers: dict, files: dict[str, bytes]) -> str:
    """在用户输出目录写入产物文件并构造一个 SUCCESS 任务。"""
    settings = get_settings()
    out_dir = settings.output_dir / str(headers["user_id"])
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, data in files.items():
        (out_dir / name).write_bytes(data)

    tm = get_task_manager()
    task_id = generate_task_id()
    tm.create(task_id, ConversionType.PNG_TO_PDF, user_id=headers["user_id"])
    tm.append_output(task_id, list(files.keys())[0])
    tm.update_status(task_id, TaskStatus.SUCCESS)
    return task_id


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["姓名", "分数"])
    ws.append(["张三", 95])
    ws.append(["李四", 88])
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Hello 预览测试")
    doc.save(buf)
    return buf.getvalue()


# =========================================================================== #
# 鉴权 / 归属
# =========================================================================== #
def test_preview_unauthorized_401(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    task_id = _make_task_with_files(
        client, headers, {"out.pdf": b"%PDF-1.4 dummy"}
    )
    resp = client.get(f"/api/v1/tasks/{task_id}/preview/out.pdf")
    assert resp.status_code == 401, resp.text


def test_preview_other_user_404(client) -> None:
    alice = _register_and_login(client, "alice", "alice123")
    bob = _register_and_login(client, "bob", "bob12345")
    task_id = _make_task_with_files(
        client, alice, {"out.pdf": b"%PDF-1.4 dummy"}
    )
    resp = client.get(
        f"/api/v1/tasks/{task_id}/preview/out.pdf",
        headers={"Authorization": "Bearer " + bob["Authorization"].split(" ")[1]},
    )
    assert resp.status_code == 404, resp.text


def test_preview_filename_not_in_task_403(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    task_id = _make_task_with_files(client, headers, {"out.pdf": b"%PDF-1.4 dummy"})
    # 文件存在但不是任务的 output_files
    resp = client.get(
        f"/api/v1/tasks/{task_id}/preview/out.pdf",
        headers={"Authorization": headers["Authorization"]},
    )
    assert resp.status_code == 200, resp.text  # 正确场景
    other = _make_task_with_files(client, headers, {"other.pdf": b"x"})
    resp = client.get(
        f"/api/v1/tasks/{other}/preview/out.pdf",
        headers={"Authorization": headers["Authorization"]},
    )
    assert resp.status_code == 403, resp.text


# =========================================================================== #
# PDF / 图片 inline
# =========================================================================== #
def test_preview_pdf_inline(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    task_id = _make_task_with_files(
        client, headers, {"out.pdf": b"%PDF-1.4 dummy content"}
    )
    resp = client.get(f"/api/v1/tasks/{task_id}/preview/out.pdf", headers={"Authorization": headers["Authorization"]})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/pdf"), resp.headers
    disposition = resp.headers.get("content-disposition", "")
    assert "inline" in disposition, disposition
    assert resp.content == b"%PDF-1.4 dummy content"


def test_preview_image_inline(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    png = b"\x89PNG\r\n\x1a\n" + b"dummy"
    task_id = _make_task_with_files(client, headers, {"img.png": png})
    resp = client.get(f"/api/v1/tasks/{task_id}/preview/img.png", headers={"Authorization": headers["Authorization"]})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("image/png"), resp.headers
    assert "inline" in resp.headers.get("content-disposition", "")


# =========================================================================== #
# XLSX / DOCX -> HTML
# =========================================================================== #
def test_preview_xlsx_html(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    task_id = _make_task_with_files(client, headers, {"out.xlsx": _xlsx_bytes()})
    resp = client.get(f"/api/v1/tasks/{task_id}/preview/out.xlsx", headers={"Authorization": headers["Authorization"]})
    assert resp.status_code == 200, resp.text
    assert "text/html" in resp.headers["content-type"], resp.headers
    body = resp.text
    assert "<table" in body, "HTML 应包含表格"
    assert "张三" in body, "HTML 应包含单元格内容"
    assert "Sheet1" in body, "HTML 应包含 Sheet 名"


def test_preview_docx_html(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    task_id = _make_task_with_files(client, headers, {"out.docx": _docx_bytes()})
    resp = client.get(f"/api/v1/tasks/{task_id}/preview/out.docx", headers={"Authorization": headers["Authorization"]})
    assert resp.status_code == 200, resp.text
    assert "text/html" in resp.headers["content-type"], resp.headers
    assert "Hello" in resp.text, "HTML 应包含段落文本"


# =========================================================================== #
# ZIP 不支持预览
# =========================================================================== #
def test_preview_zip_400(client) -> None:
    headers = _register_and_login(client, "alice", "alice123")
    task_id = _make_task_with_files(client, headers, {"batch.zip": b"PKdummy"})
    resp = client.get(f"/api/v1/tasks/{task_id}/preview/batch.zip", headers={"Authorization": headers["Authorization"]})
    assert resp.status_code == 400, resp.text
